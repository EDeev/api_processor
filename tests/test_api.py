import io
import os
import shutil
import subprocess
from concurrent import futures

import docx
import grpc
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from fpdf import FPDF
from rest_framework.test import APIClient

from api_app.grpc_client import client as grpc_client
from api_app.services.scan import extract_text_tables

import text_service_pb2
import text_service_pb2_grpc

pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def media(tmp_path, settings):
    settings.MEDIA_ROOT = str(tmp_path)


def make_docx():
    d = docx.Document()
    d.add_paragraph("Привет <script>alert(1)</script> & мир")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "a", "b"
    t.cell(1, 0).text, t.cell(1, 1).text = "1", "<2>"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pdf():
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="Hello <b>world</b> & co")
    return bytes(pdf.output())


def test_docx_text_is_escaped(tmp_path):
    path = tmp_path / "a.DOCX"
    path.write_bytes(make_docx())
    html = extract_text_tables(str(path))
    assert "<p>Привет &lt;script&gt;alert(1)&lt;/script&gt; &amp; мир</p>" in html
    assert "<pre>a,b\r\n1,&lt;2&gt;\r\n</pre>" in html


def test_pdf_upper_extension_and_escaping(tmp_path):
    path = tmp_path / "a.PDF"
    path.write_bytes(make_pdf())
    assert "<p>Hello &lt;b&gt;world&lt;/b&gt; &amp; co</p>" in extract_text_tables(str(path))


def test_document_endpoint():
    client = APIClient()
    resp = client.post("/api/document-to-text/",
                       {"document": SimpleUploadedFile("r.docx", make_docx())}, format="multipart")
    assert resp.status_code == 200
    assert "&lt;script&gt;" in resp.json()["text"] and resp.json()["grpc_response"] is None


def test_validation_and_size_limit(settings):
    client = APIClient()
    assert client.post("/api/document-to-text/", {}, format="multipart").status_code == 400
    bad = client.post("/api/document-to-text/", {"document": SimpleUploadedFile("a.txt", b"x")}, format="multipart")
    assert bad.status_code == 400
    settings.MAX_UPLOAD_SIZE = 10
    big = client.post("/api/document-to-text/", {"document": SimpleUploadedFile("a.pdf", b"x" * 100)},
                      format="multipart")
    assert big.status_code == 413


@override_settings(API_TOKEN="secret")
def test_api_token():
    client = APIClient()
    upload = {"document": SimpleUploadedFile("r.docx", make_docx())}
    assert client.post("/api/document-to-text/", upload, format="multipart").status_code == 403
    client.credentials(HTTP_AUTHORIZATION="Bearer secret")
    upload = {"document": SimpleUploadedFile("r.docx", make_docx())}
    assert client.post("/api/document-to-text/", upload, format="multipart").status_code == 200


class Upper(text_service_pb2_grpc.TextProcessorServicer):
    def ProcessText(self, request, context):
        return text_service_pb2.TextResponse(processed_text=request.text.upper(), success=True)


def test_grpc_real_server(settings):
    server = grpc.server(futures.ThreadPoolExecutor(max_workers=1))
    text_service_pb2_grpc.add_TextProcessorServicer_to_server(Upper(), server)
    port = server.add_insecure_port("127.0.0.1:0")
    server.start()
    try:
        settings.GRPC_SERVER = f"127.0.0.1:{port}"
        assert grpc_client.send_to_grpc_server("привет") == {
            "processed_text": "ПРИВЕТ", "success": True, "error": None}
    finally:
        server.stop(None)


def test_grpc_unavailable(settings):
    settings.GRPC_SERVER = "127.0.0.1:1"
    settings.GRPC_TIMEOUT = 2
    result = grpc_client.send_to_grpc_server("x")
    assert result["success"] is False and result["error"].startswith("gRPC")


@pytest.mark.skipif(not shutil.which("ffmpeg"), reason="нужен ffmpeg")
def test_audio_endpoint_converts_ogg(tmp_path, settings):
    if not os.path.isdir(settings.VOSK_MODEL_PATH):
        pytest.skip("нет модели Vosk")
    ogg = tmp_path / "tone.ogg"
    subprocess.run(["ffmpeg", "-loglevel", "error", "-f", "lavfi", "-i", "sine=frequency=440:duration=2",
                    "-c:a", "libvorbis", str(ogg)], check=True)
    resp = APIClient().post("/api/audio-to-text/",
                            {"audio": SimpleUploadedFile("tone.ogg", ogg.read_bytes())}, format="multipart")
    assert resp.status_code == 200 and isinstance(resp.json()["text"], str)


SAMPLE = os.path.join(os.path.dirname(os.path.dirname(__file__)), "examples", "sample.ogg")


@pytest.mark.skipif(not shutil.which("ffmpeg"), reason="нужен ffmpeg")
def test_sample_recognition_keeps_all_phrases(settings):
    if not os.path.isdir(settings.VOSK_MODEL_PATH):
        pytest.skip("нет модели Vosk")
    from api_app.services.vosk_recognizer import recognize_speech

    text = recognize_speech(SAMPLE)
    # в записи несколько фраз: раньше оставалась только последняя
    assert "проверка" in text and "десять" in text
